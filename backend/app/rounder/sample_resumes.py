"""Sample resumes built with python-docx, used by tests, `make rounder-eval` and the E2E test.

The people and companies are made up. Each builder covers a layout the template reader has to
handle: Word heading styles, bold all-caps headings with typed bullets, experience laid out in a
table, and a resume with no experience section.

    uv run python -m app.rounder.sample_resumes OUT_DIR   # writes every sample as a .docx
"""

import io
import sys
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

Entry = tuple[str, str, list[str]]  # (role line, dates line, bullets)

ENTRIES: list[Entry] = [
    (
        "Senior Software Engineer, Northwind Analytics",
        "Mar 2021 \u2013 Present",
        [
            "Built and maintained REST APIs in Python serving the analytics dashboard.",
            "Moved nightly reporting jobs from cron scripts to a queue-based worker system.",
            "Reviewed code and mentored two junior engineers on testing practices.",
            "Cut page load time of the main dashboard by 40% by caching query results.",
        ],
    ),
    (
        "Software Engineer, Fabrikam Logistics",
        "Jun 2018 \u2013 Feb 2021",
        [
            "Developed shipment tracking features in a Django web application.",
            "Wrote SQL reports for the operations team and tuned slow queries.",
            "Set up automated tests and a build pipeline for the main service.",
        ],
    ),
    (
        "Junior Developer, Contoso Retail",
        "Aug 2016 \u2013 May 2018",
        [
            "Maintained the online store's checkout pages.",
            "Fixed bugs reported by customer support within agreed response times.",
        ],
    ),
]

EXTRA_ENTRIES: list[Entry] = [
    (
        "Web Developer, Tailspin Media",
        "Jan 2014 \u2013 Jul 2016",
        [
            "Built marketing sites for clients from designers' mock-ups.",
            "Added analytics tracking and A/B tests to landing pages.",
            "Managed releases and hosting for twelve client sites.",
        ],
    ),
    (
        "IT Support Technician, Wingtip Toys",
        "Sep 2011 \u2013 Dec 2013",
        [
            "Supported 150 staff with hardware, software and network issues.",
            "Wrote scripts to automate laptop setup for new starters.",
            "Kept the asset inventory and software licences up to date.",
        ],
    ),
    (
        "Freelance Developer",
        "Jan 2010 \u2013 Aug 2011",
        [
            "Built small business websites with PHP and MySQL.",
            "Moved three clients from shared hosting to managed virtual servers.",
            "Trained client staff to update their own site content.",
        ],
    ),
]

SUMMARY = (
    "Software engineer with eight years of experience building web applications and data "
    "tools. Enjoys making slow systems fast and helping teams ship with confidence."
)
EDUCATION = "BSc Computer Science, University of Example, 2016"
SKILLS_LINE = "Python, Django, SQL, Git, Linux, Docker"
PROJECTS = [
    "Open-source command-line tool for comparing CSV files, with 300 stars on GitHub.",
    "Volunteer-built booking site for a community sports club, used weekly by 80 members.",
    "Home weather station that logs readings to a small time-series database.",
    "Browser extension that hides distracting elements on news sites.",
    "Talk at a local meetup on testing data pipelines.",
    "Workshop for new programmers on version control with Git.",
]


def _base(long: bool = False) -> DocxDocument:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Liberation Serif"
    style.font.size = Pt(12 if long else 11)
    if long:
        style.paragraph_format.space_after = Pt(8)
    return doc


def _contact(doc: DocxDocument) -> None:
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("Alex Rivera")
    run.bold = True
    run.font.size = Pt(18)
    contact = doc.add_paragraph("alex.rivera@example.com · +1 555 0100 · Springfield")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _entries(long: bool) -> list[Entry]:
    return ENTRIES + EXTRA_ENTRIES if long else ENTRIES


def styled(long: bool = False) -> bytes:
    """Word heading styles; entry headers as Heading 2; bullets in the "List Bullet" style."""
    doc = _base(long)
    _contact(doc)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(SUMMARY)
    doc.add_heading("Experience", level=1)
    for role, dates, bullets in _entries(long):
        doc.add_heading(role, level=2)
        doc.add_paragraph(dates).runs[0].italic = True
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
    if long:
        # Bullets after the experience section: they must never be rewritten.
        doc.add_heading("Projects", level=1)
        for project in PROJECTS:
            doc.add_paragraph(project, style="List Bullet")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(EDUCATION)
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(SKILLS_LINE)
    return _save(doc)


def _caps_heading(doc: DocxDocument, text: str) -> None:
    run = doc.add_paragraph().add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)


def caps(long: bool = False) -> bytes:
    """No heading styles: bold all-caps headings, bold role lines and typed "•" bullets."""
    doc = _base()
    _contact(doc)
    _caps_heading(doc, "Profile")
    doc.add_paragraph(SUMMARY)
    _caps_heading(doc, "Professional Experience")
    for role, dates, bullets in _entries(long):
        header = doc.add_paragraph()
        header.add_run(role).bold = True
        header.add_run(f"\t{dates}")
        for bullet in bullets:
            p = doc.add_paragraph(f"•\t{bullet}")
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
    _caps_heading(doc, "Education")
    doc.add_paragraph(EDUCATION)
    return _save(doc)


def table(long: bool = False) -> bytes:
    """Experience laid out in a two-column table: dates on the left, role and bullets right."""
    doc = _base()
    _contact(doc)
    doc.add_heading("Work Experience", level=1)
    grid = doc.add_table(rows=0, cols=2)
    for role, dates, bullets in _entries(long):
        cells = grid.add_row().cells
        cells[0].paragraphs[0].add_run(dates).italic = True
        cells[1].paragraphs[0].add_run(role).bold = True
        for bullet in bullets:
            cells[1].add_paragraph(bullet, style="List Bullet")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(EDUCATION)
    return _save(doc)


def no_experience() -> bytes:
    """Sections a person might use instead of "Experience": none of them should be guessed."""
    doc = _base()
    _contact(doc)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(SUMMARY)
    doc.add_heading("Projects", level=1)
    for _, _, bullets in ENTRIES[:2]:
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(EDUCATION)
    return _save(doc)


def _save(doc: DocxDocument) -> bytes:
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


SAMPLES: dict[str, Callable[[], bytes]] = {
    "styled-1-page": styled,
    "styled-2-pages": lambda: styled(long=True),
    "caps-headings": caps,
    "table-layout": table,
    "no-experience": no_experience,
}


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__)
        return 2
    out_dir = Path(argv[0])
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, build in SAMPLES.items():
        (out_dir / f"{name}.docx").write_bytes(build())
        print(out_dir / f"{name}.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
